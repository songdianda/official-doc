#!/usr/bin/env python3
"""生成前后内容验收：比对段落文字、表格与图片，优先拦住漏字、漏表、漏图。

用法：
    python3 scripts/verify_output.py 源文件.docx 成品.docx [提取结果.json]

这个脚本**不检查版式**。版式（字体、行距、页码位置）要另外逐页看。
它只回答一个问题：源文档里的东西，成品里还在不在。

四道检查，各自独立、互不替代
  1. 原文覆盖    **直接从源文档读**，不经过提取结果。
                 源里每一段（含表格格、页眉页脚）都必须在成品里找得到，
                 找不到就是丢了。这一道专门防「提取 JSON 和成品同时漏掉同一段」
                 ——那种情况若只拿 JSON 当基准，会得出"一致"的假结论。
  2. 反向检查    成品里每一段也必须能在源里配对，配对后**多出来的字符**就是增字。
                 只查"源有成品无"是不够的：给段落加个「二、」前缀，
                 正向检查照样判一致——加字只能靠反向比对抓出来。
  3. 表格逐格    按 (行, 列) 坐标比对，不是比内容集合。
                 集合比对抓不到"两格内容互换"和"重复值多一个少一个"。
                 合并单元格先按 w:gridSpan / w:vMerge 还原成逻辑格再比。
  4. 图片        按部位分别计数：正文 / 页眉页脚 / 脚注尾注批注。
                 只看 body 会把页眉里的图片数成 0，页眉图丢了也报"一致"。

严格保真（fidelity=strict，默认）
  原文一个字不增不减。反向检查发现的增字、以及源文档没有却出现在成品里的
  页码，都判为不通过。确属有意补充时用 --fidelity standard，增字降级为提示。

有意调整的部分（不是缺陷，会在报告里单独说明）
  - 落款署名 / 成文日期被剥离到 signature / date，成品中位置改变但内容仍在
  - 版头整段原样复制（文本框内的字需目视确认）
  - 落款前后用于「顶下去」的空段落被清理
  - standard 模式下 Word 自动编号被补回成文字（源文件里它不是字符）
"""

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document                      # noqa: E402
from docx.oxml.ns import qn                    # noqa: E402

# VML 命名空间不在 python-docx 的 nsmap 里，qn('v:line') 会 KeyError
VML_NS = 'urn:schemas-microsoft-com:vml'


def _vml(tag):
    return '{%s}%s' % (VML_NS, tag)

# 只复用纯 XML 工具函数，不复用提取逻辑——验收必须独立于被验收的对象
from inspect_docx import extract_table         # noqa: E402


def _norm(s):
    """把连续空白折叠成一个空格，再 trim，用于精确比较。

    只折叠空白，**不做子串匹配**——完整文字是默认口径。
    「检测结果不合格，禁止投入使用。」和「检测结果」必须判为不同，
    不能因为后者是前者子串就放行（那是截断/少字，要拦截）。
    """
    return re.sub(r'\s+', ' ', (s or '')).strip()


def _all_text(el):
    """元素内全部 w:t 文字（含文本框内的）。

    文本框里的 w:t 嵌在 w:pict 深处，python-docx 的 p.text 用相对 xpath
    只取直接子节点，取不到——红头是文本框时会读成空段，必须自己遍历。
    """
    return ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()


# ── 图片计数 ─────────────────────────────────────────────

def _count_images(el):
    """XML 元素里的**位图**引用数。

    只数真正引用了图片文件的元素：DrawingML 的 a:blip 带 r:embed，
    或老式 VML 的 v:imagedata 带 r:id。红头分隔线、版记线这类 v:line /
    v:shape 是矢量图形，没有位图文件，不能算图片——把它们数进来会报出
    「正文图片少 4 个」这种假警报（实测样本：源 6 / 成品 2，实际一张
    位图都没丢，少的全是 VML 线条）。
    """
    n = 0
    for blip in el.iter(qn('a:blip')):
        if blip.get(qn('r:embed')) or blip.get(qn('r:link')):
            n += 1
    for im in el.iter(_vml('imagedata')):
        if im.get(qn('r:id')):
            n += 1
    return n


def _is_page_number(el):
    """元素里是否含 PAGE 域（页码）。

    页脚里常同时存在域代码和上次渲染的缓存数字（如「8」），把缓存值当
    正文比会误报「页眉页脚文字未进入成品」——成品有同一个域，只是缓存
    值不同。页码由专门的那一道检查负责。
    """
    for t in el.iter(qn('w:instrText')):
        if 'PAGE' in (t.text or '').upper():
            return True
    return False


def _header_footer_text(el):
    """提取页眉/页脚元素里的**非页码**文字。

    页码域（w:instrText 含 PAGE）及其缓存数字不是内容，页码单独比对；
    但页脚里除了页码域之外的文字（如「内部资料，请勿外传」「机密」）
    是真实内容，必须参与验收——不能因为页脚含 PAGE 域就整段跳过。

    做法：按 run 遍历。一个 run 若含域标记（w:fldChar / w:instrText）
    或其内容看起来就是 PAGE 域渲染出的纯数字缓存，则跳过；否则保留文字。
    """
    def _is_field_run(r):
        return any(True for _ in r.iter(qn('w:fldChar'))) or \
               any(True for _ in r.iter(qn('w:instrText')))

    parts = []
    for p in el.iter(qn('w:p')):
        has_page = any('PAGE' in (t.text or '').upper()
                       for t in p.iter(qn('w:instrText')))
        for r in p.iter(qn('w:r')):
            if _is_field_run(r):
                continue            # 域代码 / 域开关本身不是正文文字
            for t in r.iter(qn('w:t')):
                txt = (t.text or '').strip()
                if not txt:
                    continue
                if has_page and re.fullmatch(r'[0-9]{1,4}', txt):
                    continue        # 页码域的缓存数字，页码单独比对
                parts.append(t.text)
    return re.sub(r'\s+', ' ', ''.join(parts)).strip()


def _count_shapes(el):
    """矢量图形数（分隔线、版记线、形状框），与位图分开报。"""
    tags = tuple(_vml(t) for t in
                 ('line', 'shape', 'rect', 'oval', 'roundrect'))
    return sum(1 for node in el.iter() if node.tag in tags)


def _classify(doc):
    """把包内部件按部位分类。

    按 partname 分类而不按 section.header / section.footer 取：
    首页页眉、偶数页页眉不一定挂在默认属性上，逐个属性取会漏；
    直接扫部件则只要是包里的页眉页脚就都在。
    """
    zones = {'body': [], 'header_footer': [], 'other': []}
    for part in doc.part.package.parts:
        name = str(part.partname)
        if not name.startswith('/word/') or not hasattr(part, 'element'):
            continue
        base = name.rsplit('/', 1)[-1]
        if base == 'document.xml':
            zones['body'].append(part.element)
        elif base.startswith('header') or base.startswith('footer'):
            zones['header_footer'].append(part.element)
        elif base in ('footnotes.xml', 'endnotes.xml', 'comments.xml'):
            zones['other'].append(part.element)
    return zones


def image_stats(path):
    """按部位统计图片：正文 / 页眉页脚 / 脚注尾注批注。

    只数 body 是不够的：红头图片常在页眉里，页眉图丢了却报「双方 0 张」，
    是最容易被忽略的一类内容丢失。
    """
    doc = Document(path)
    zones = _classify(doc)
    stats = {k: sum(_count_images(el) for el in v) for k, v in zones.items()}
    stats['shapes'] = sum(_count_shapes(el) for el in zones['body'])
    # 去重后的图片文件数（同一张图在多个页眉被引用时只算一份）
    stats['files'] = len({str(p.partname) for p in doc.part.package.parts
                          if '/media/' in str(p.partname)})
    return stats


def numbering_marks(path):
    """统计正文里带自动编号（w:numPr）的段落，返回 [(文本, numId, ilvl)]。

    自动编号不是字符，`p.text` 读不到——必须看 w:numPr 才能知道这段
    有没有编号、编到哪一级。验收据此判断编号是否丢失。
    """
    doc = Document(path)
    marks = []
    for p in doc.paragraphs:
        pPr = p._element.pPr
        if pPr is None:
            continue
        numPr = pPr.find(qn('w:numPr'))
        if numPr is None:
            continue
        ilvl = numPr.find(qn('w:ilvl'))
        numid = numPr.find(qn('w:numId'))
        marks.append({
            'text': _all_text(p._element),
            'numId': numid.get(qn('w:val')) if numid is not None else None,
            'ilvl': ilvl.get(qn('w:val')) if ilvl is not None else None,
        })
    return marks


# ── 直接从文档读原文（不依赖任何提取结果）─────────────────

def raw_source(path):
    """源文档里"应该被带走的全部文字"，逐条带位置。

    返回 list of {'text', 'where', 'zone'}，zone ∈ 正文 / 表格 / 页眉页脚。
    这一份清单**不经过** inspect_docx 的角色判定与落款剥离，
    是验收的独立基准：提取环节漏掉的东西，在这里仍然看得到。
    """
    doc = Document(path)
    items = []

    for i, p in enumerate(doc.paragraphs):
        text = _all_text(p._element)
        if text:
            items.append({'text': text, 'where': f'正文段落[{i}]',
                          'zone': '正文'})

    for ti, tbl in enumerate(doc.tables):
        rows, _spans, _ncol = extract_table(tbl)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                if val.strip():
                    items.append({'text': val.strip(),
                                  'where': f'表{ti + 1}({r + 1}行{c + 1}列)',
                                  'zone': '表格'})

    for el in _classify(doc)['header_footer']:
        text = _header_footer_text(el)
        # 页眉页脚里的文字（页码域除外）是真实内容，参与必须通过的验收
        if text:
            items.append({'text': text, 'where': '页眉/页脚',
                          'zone': '页眉页脚'})
    for el in _classify(doc)['other']:
        text = _all_text(el)
        if text:
            items.append({'text': text, 'where': '脚注/尾注/批注',
                          'zone': '其他'})
    return items


def target_texts(path):
    """成品里全部可比对文字：正文段落 + 表格格 + 页眉页脚。"""
    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        text = _all_text(p._element)
        if text:
            out.append(text)
    for tbl in doc.tables:
        rows, _spans, _ncol = extract_table(tbl)
        for row in rows:
            out.extend(v.strip() for v in row if v.strip())
    zones = _classify(doc)
    for el in zones['header_footer']:
        text = _header_footer_text(el)
        if text:
            out.append(text)
    for el in zones['other']:
        text = _all_text(el)
        if text:
            out.append(text)
    return out


def match_one(text, pool):
    """在 pool 中找一条**精确相等**（空白折叠后）的记录。

    默认完整文字比较，不再互为子串放行。命中后把该条从 pool 消耗掉：
    源里有两段相同文字、成品只有一段时，第二段找不到配对，从而暴露
    「少了一段」——只比集合发现不了，靠消耗来保留重复段落的次数。
    """
    n = _norm(text)
    if not n:
        return True
    for i, cand in enumerate(pool):
        if _norm(cand) == n:
            pool.pop(i)
            return True
    return False


def _diff_detail(src, dst):
    """给出源与成品的文字差异描述（少字/多字/替换）。

    返回 (kind, detail)：
      kind ∈ 'equal' / 'missing'(少字) / 'extra'(多字) / 'replaced'(替换)
    detail 为可读描述，含原文与成品片段。
    """
    s, d = _norm(src), _norm(dst)
    if s == d:
        return 'equal', ''
    if not s:
        return 'extra', f'源为空，成品多出「{d[:40]}」'
    if not d:
        return 'missing', f'源「{s[:40]}」在成品中缺失'
    if s in d:
        # 成品比源多：定位多出的前后缀
        idx = d.find(s)
        extra = (d[:idx] + d[idx + len(s):]).strip()
        return 'extra', f'成品多出「{extra[:40]}」'
    if d in s:
        # 成品比源少：源的一部分被截断
        idx = s.find(d)
        missing = (s[:idx] + s[idx + len(d):]).strip()
        return 'missing', f'缺失「{missing[:40]}」'
    # 中间替换：用 difflib 找第一个不同片段
    import difflib
    sm = difflib.SequenceMatcher(None, s, d)
    first = None
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag != 'equal':
            first = (tag, s[i1:i2], d[j1:j2])
            break
    if first:
        tag, a, b = first
        if tag == 'delete':
            return 'missing', f'缺失「{a[:40]}」（源「{s[:30]}」→ 成品「{d[:30]}」）'
        if tag == 'insert':
            return 'extra', f'多出「{b[:40]}」（源「{s[:30]}」→ 成品「{d[:30]}」）'
        return 'replaced', f'「{a[:20]}」→「{b[:20]}」（源「{s[:30]}」→ 成品「{d[:30]}」）'
    return 'replaced', f'源「{s[:40]}」→ 成品「{d[:40]}」'


def match_extra(text, pool):
    """反向配对：成品一段文字在源 pool 里找精确对应。

    精确匹配不上就给出可读差异描述（少字/多字/替换），不再把「成品是源
    子串」判为 0 增字——少字要在这里被拦住，而不是推给正向检查。
    找不到任何相近源段时返回整段原文（视为整段新增）。
    """
    n = _norm(text)
    if not n:
        return ''
    for i, cand in enumerate(pool):
        if _norm(cand) == n:
            pool.pop(i)
            return ''
    # 没精确匹配上：找最接近的一条源段给出可读差异
    import difflib
    best_i, best_ratio = None, 0.0
    for i, cand in enumerate(pool):
        ratio = difflib.SequenceMatcher(None, n, _norm(cand)).ratio()
        if ratio > best_ratio:
            best_i, best_ratio = i, ratio
    if best_i is not None and best_ratio >= 0.5:
        cand = pool.pop(best_i)
        kind, detail = _diff_detail(cand, text)
        return detail if kind != 'equal' else n
    return n                  # 源里完全没有 → 整段都是新增


def body_texts(path):
    """成品正文段落文字（不含表格、不含页眉页脚），逐条带位置。

    页眉页脚必须排除：页码在页脚里，算进来会整段被当成增字。
    """
    doc = Document(path)
    out = []
    for i, p in enumerate(doc.paragraphs):
        text = _all_text(p._element)
        if text:
            out.append({'text': text, 'where': f'成品正文段落[{i}]'})
    return out


def source_body_ordered(path):
    """源文档正文段落（zone=正文）的有序文字序列，用于顺序校验。

    返回 [text, ...]，保留出现次序与重复次数。
    """
    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        text = _all_text(p._element)
        if text:
            out.append(text)
    return out


def check_order(src_seq, dst_seq):
    """检测正文段落顺序是否变化（未经允许的段落顺序变化）。

    思路：先看源与成品正文段落的内容**多重集合**（含重复）是否一致。
    一致但**序列**不同 → 内容没丢没增，但发生了重排，必须报告。
    集合不一致 → 交给原文覆盖 / 反向检查负责，这里不重复报。

    返回 (是否一致, 差异描述列表)。
    """
    from collections import Counter
    sn = [_norm(x) for x in src_seq]
    dn = [_norm(x) for x in dst_seq]
    if Counter(sn) != Counter(dn):
        # 内容本身有增删，顺序问题由覆盖/反向检查覆盖，此处不重复报
        return True, []
    if sn == dn:
        return True, []
    # 集合一致但序列不同 → 重排。逐位找第一个不一致的位置
    reorders = []
    for i, (a, b) in enumerate(zip(sn, dn)):
        if a != b:
            reorders.append(
                f'第{i + 1}段起顺序变化：源「{src_seq[i][:25]}」'
                f'↔ 成品「{dst_seq[i][:25]}」')
            break
    if not reorders and len(sn) != len(dn):
        reorders.append('段落数量一致但顺序不同')
    return (not reorders), reorders


def page_number_marks(path):
    """页脚里是否有页码域（w:instrText 含 PAGE）。

    只看文字不行：页码是域，未更新时取不到数字，但域代码一直在。
    """
    doc = Document(path)
    marks = []
    for el in _classify(doc)['header_footer']:
        instr = ''.join(t.text or '' for t in el.iter(qn('w:instrText')))
        if 'PAGE' in instr.upper():
            marks.append(instr.strip() or '(空域)')
    return marks


# ── 表格：按坐标逐格比对 ──────────────────────────────────

def table_grid(path):
    """读文档的表格，还原合并后返回 [(行数, 列数, {(r,c): 文本})]。"""
    doc = Document(path)
    grids = []
    for tbl in doc.tables:
        rows, spans, ncol = extract_table(tbl)
        cells = {}
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cells[(r, c)] = val.strip()
        grids.append((len(rows), ncol, cells))
    return grids


def compare_tables(src_grids, dst_grids):
    """逐格比对，返回 (是否一致, 差异明细)。"""
    diffs = []
    if len(src_grids) != len(dst_grids):
        diffs.append(f'表格数量不一致：源 {len(src_grids)} 个，'
                     f'成品 {len(dst_grids)} 个')
    for i, (s, d) in enumerate(zip(src_grids, dst_grids)):
        srow, scol, scells = s
        drow, dcol, dcells = d
        if (srow, scol) != (drow, dcol):
            diffs.append(f'表{i + 1} 结构不同：源 {srow}行×{scol}列，'
                         f'成品 {drow}行×{dcol}列')
        keys = sorted(set(scells) | set(dcells))
        for (r, c) in keys:
            a, b = scells.get((r, c), ''), dcells.get((r, c), '')
            if a != b:
                diffs.append(f'表{i + 1}({r + 1}行{c + 1}列)：'
                             f'源「{a[:30]}」→ 成品「{b[:30]}」')
    return (not diffs), diffs


def json_items(meta):
    """提取结果（JSON）里的文字清单，用于判断漏在哪个环节。"""
    items = []
    if not meta:
        return items
    if meta.get('title'):
        items.append(str(meta['title']).strip())
    for it in meta.get('body', []):
        txt = (it.get('text') or '').strip() if isinstance(it, dict) else ''
        if txt:
            items.append(txt)
    for key in ('signature', 'date'):
        if meta.get(key):
            items.append(str(meta[key]).strip())
    for t in (meta.get('tables') or []):
        for row in t.get('rows', []):
            items.extend(str(v).strip() for v in row if str(v).strip())
    for a in (meta.get('attachments') or []):
        items.append(str(a).strip())
    return [x for x in items if x]


def main():
    ap = argparse.ArgumentParser(
        description='比对源文档与成品的内容，检查是否有漏字、漏表、漏图')
    ap.add_argument('source', help='源 .docx')
    ap.add_argument('target', help='生成的 .docx')
    ap.add_argument('meta', nargs='?',
                    help='inspect_docx.py 输出的 JSON（可选，用于定位漏在哪个环节）')
    ap.add_argument('--fidelity', choices=('strict', 'standard'), default=None,
                    help='strict（默认）：增字与新增页码判不通过；'
                         'standard：降级为提示')
    args = ap.parse_args()

    for p in (args.source, args.target):
        if not os.path.exists(p):
            print(f'文件不存在: {p}', file=sys.stderr)
            return 1

    meta = None
    if args.meta:
        import json
        if not os.path.exists(args.meta):
            print(f'提取结果不存在: {args.meta}', file=sys.stderr)
            return 1
        meta = json.load(open(args.meta, encoding='utf-8'))

    # 保真模式：优先命令行，其次提取结果，最后默认 strict
    fidelity = args.fidelity or (meta or {}).get('fidelity') or 'strict'
    strict = fidelity == 'strict'

    ok = True
    warns = []
    partial = []          # 部分完成/待确认（不支持内容、无法判定的内容）
    print('──────── 内容验收报告 ────────')
    print(f'源文件：{os.path.basename(args.source)}')
    print(f'成品　：{os.path.basename(args.target)}')

    # ── 1. 原文覆盖（基准来自源文档本身）────────────────
    src_items = raw_source(args.source)
    pool = target_texts(args.target)
    jpool = json_items(meta)
    missing = []
    for it in src_items:
        if not match_one(it['text'], pool):
            missing.append(it)
    print(f'\n【原文覆盖】源文档 {len(src_items)} 条文字，逐条在成品中查找')
    if missing:
        # 正文、表格、页眉页脚都是"必须保留"的内容，丢失判不通过；
        # 只有"其他"（脚注/尾注/批注）属本技能不处理的内容 → 部分完成
        hard = [m for m in missing if m['zone'] in ('正文', '表格', '页眉页脚')]
        soft = [m for m in missing if m['zone'] == '其他']
        if hard:
            ok = False
            print(f'  成品中找不到 {len(hard)} 条（内容丢失）：')
            for m in hard[:20]:
                mark = ''
                if meta and not match_one(m['text'], jpool):
                    mark = '　[提取结果里也没有 → 漏在提取环节，不是排版环节]'
                print(f'    × {m["where"]}  「{m["text"][:60]}」{mark}')
            if len(hard) > 20:
                print(f'    …… 另有 {len(hard) - 20} 条')
        if not hard:
            print('  正文、表格、页眉页脚文字均可在成品中找到'
                  '（脚注/尾注/批注等缺失见文末提示）')
        if soft:
            partial.append(
                f'脚注/尾注/批注文字 {len(soft)} 条未进入成品'
                f'（本技能不处理这些内容，需手工补回）：'
                + '、'.join(m['text'][:20] for m in soft[:5]))
    else:
        print('  源文档全部文字均可在成品中找到')

    # ── 2. 反向检查：成品里有没有源文档没有的字 ──────────
    dst_body = body_texts(args.target)
    rpool = [it['text'] for it in src_items if it['zone'] in ('正文', '表格')]
    extras = []
    for it in dst_body:
        more = match_extra(it['text'], rpool)
        if more:
            extras.append((it, more))
    print(f'\n【反向检查】成品 {len(dst_body)} 段正文，逐段在源文档中配对'
          f'（保真模式：{fidelity}）')
    if not extras:
        print('  成品正文没有源文档以外的字符')
    else:
        head = ('  发现 %d 段与源文档不符（少字/多字/替换）：' % len(extras))
        if strict:
            ok = False
            print(head)
        else:
            warns.append(head.strip())
            print(head + '（standard 模式，记为提示）')
        for it, more in extras[:20]:
            print(f'    × {it["where"]}  {more}｜成品整段：{it["text"][:50]}')
        if len(extras) > 20:
            print(f'    …… 另有 {len(extras) - 20} 段')

    # ── 2b. 段落顺序：未经允许的重排 ────────────────────
    src_seq = source_body_ordered(args.source)
    dst_seq = [it['text'] for it in dst_body]
    order_ok, reorders = check_order(src_seq, dst_seq)
    if reorders:
        ok = False
        print(f'\n【段落顺序】发现 {len(reorders)} 处未经允许的段落顺序变化：')
        for r in reorders[:20]:
            print(f'    × {r}')
        if len(reorders) > 20:
            print(f'    …… 另有 {len(reorders) - 20} 处')

    # ── 页码：源文档没有就不该有 ────────────────────────
    src_pn = page_number_marks(args.source)
    dst_pn = page_number_marks(args.target)
    if dst_pn and not src_pn:
        msg = ('源文档没有页码，成品页脚出现 %d 处页码域 —— 属于新增内容'
               % len(dst_pn))
        if strict:
            ok = False
            print(f'\n【页码】{msg}')
        else:
            warns.append(msg)
            print(f'\n【页码】{msg}（standard 模式，记为提示）')
        print('  确认要加页码的话：这是你显式指定的，可忽略；'
              '否则检查 config 里的 page_number 是否被误设为 center / odd_even')
    elif src_pn and not dst_pn:
        ok = False
        print('\n【页码】源文档有页码，成品未保留 —— 可见内容丢失，判不通过')
    else:
        print('\n【页码】源与成品一致'
              + ('（均无页码域）' if not src_pn else '（均含页码域）'))

    # ── 3. 表格逐格（带位置）────────────────────────────
    src_grids = table_grid(args.source)
    dst_grids = table_grid(args.target)
    same, diffs = compare_tables(src_grids, dst_grids)
    total_cells = sum(len(g[2]) for g in src_grids)
    print(f'\n【表格逐格】源 {len(src_grids)} 个表 / 成品 {len(dst_grids)} 个表，'
          f'源共 {total_cells} 个格位')
    if not src_grids and not dst_grids:
        print('  双方均无表格')
    elif same:
        print(f'  逐格一致：位置与内容均相同（含合并单元格还原）')
    else:
        ok = False
        print(f'  发现 {len(diffs)} 处差异（位置不同也算差异）：')
        for d in diffs[:20]:
            print(f'    × {d}')
        if len(diffs) > 20:
            print(f'    …… 另有 {len(diffs) - 20} 处')

    # ── 3. 图片（分部位）───────────────────────────────
    si, di = image_stats(args.source), image_stats(args.target)
    print(f'\n【图片】源 正文 {si["body"]} / 页眉页脚 {si["header_footer"]}'
          f' / 其他 {si["other"]}（去重后 {si["files"]} 个图片文件）')
    print(f'　　　　成品 正文 {di["body"]} / 页眉页脚 {di["header_footer"]}'
          f' / 其他 {di["other"]}（去重后 {di["files"]} 个图片文件）')
    if si['body'] != di['body']:
        ok = False
        if si['body'] > di['body']:
            print(f'  正文图片少 {si["body"] - di["body"]} 个：'
                  f'本技能不重排正文图片，需手工补回')
        else:
            print(f'  成品正文图片比源多 {di["body"] - si["body"]} 个，请核对')
    if si['header_footer'] != di['header_footer']:
        ok = False
        if si['header_footer'] > di['header_footer']:
            print(f'  页眉页脚图片少 '
                  f'{si["header_footer"] - di["header_footer"]} 个：'
                  f'若源文档的版头在页眉里，需手工补回')
        else:
            print(f'  成品页眉页脚图片比源多 '
                  f'{di["header_footer"] - si["header_footer"]} 个，请核对')
    if si['other'] != di['other']:
        warns.append(f'脚注/尾注/批注中的图片：源 {si["other"]}、'
                     f'成品 {di["other"]}（本技能不处理这些内容）')
    if (si['body'], si['header_footer']) == (di['body'], di['header_footer']):
        print('  正文与页眉页脚图片数量一致')

    # ── 3b. 自动编号（numPr）：可见编号是否保留 ──────────
    src_num = numbering_marks(args.source)
    dst_num = numbering_marks(args.target)
    print(f'\n【自动编号】源 {len(src_num)} 段 / 成品 {len(dst_num)} 段带 Word 自动编号')
    if src_num and not dst_num:
        ok = False
        print('  源文档有自动编号，成品全部丢失 —— 可见编号与层级缺失，判不通过')
    elif len(dst_num) < len(src_num):
        ok = False
        print(f'  成品自动编号段落少于源（源 {len(src_num)}、成品 {len(dst_num)}）'
              f' —— 部分编号丢失，判不通过')
    elif src_num:
        # 逐段核对：源每段带编号的段落，成品里应有对应编号（按文本配对）
        src_texts = {m['text'] for m in src_num}
        dst_texts = {m['text'] for m in dst_num}
        lost = src_texts - dst_texts
        if lost:
            ok = False
            print(f'  以下段落的编号在成品中缺失：')
            for t in list(lost)[:10]:
                print(f'    × 「{t[:40]}」')
        else:
            print('  自动编号段落已保留（编号由 Word 渲染，非正文字符）')

    # ── 有意调整 ───────────────────────────────────────
    notes = []
    if meta:
        st = meta.get('_stripped') or []
        if st:
            notes.append('落款留在源文档原位（正文之后、附件之前）输出：'
                         + '、'.join(f"{s['field']}「{s['text']}」" for s in st)
                         + '；若成品中它被排在文末，说明落款位置丢了')
        an = meta.get('_auto_number') or []
        applied = [a for a in an if a.get('applied')]
        if applied:
            notes.append(f'Word 自动编号 {len(applied)} 处已补回为文字'
                         f'（源文件里不是字符，属 standard 模式的有意补充）')
        if meta.get('header'):
            n = len(meta['header']['indices'])
            notes.append(f'版头 {n} 段整段原样复制；'
                         f'若含文本框，其中文字已按全文 w:t 参与比对，'
                         f'框体与位置仍需目视确认')
        for u in (meta.get('_unsupported') or []):
            # 不支持的内容（图片/脚注/公式等）本技能不保留，
            # 属「部分完成／待确认」，不能算完全通过
            partial.append(
                f"源文档含「{u['kind']}」{u['count']} 处：{u['hint']}")
    notes.extend(warns)
    if notes:
        print('\n【有意调整 / 需留意】')
        for n in notes:
            print(f'  - {n}')
    if partial:
        print('\n【部分完成 / 待确认】')
        for p in partial:
            print(f'  ! {p}')

    # 三类结论 + 三档退出码：
    #   0 完全通过     —— 无内容丢失、无增字、无未处理内容
    #   1 部分完成     —— 存在本技能不处理/无法确认的内容（脚注、批注、图片等）
    #   2 不通过       —— 存在内容丢失、增字、替换、漏表、漏图等硬差异
    if not ok:
        verdict = '不通过：存在内容差异，请按上述条目核对后再交付'
        code = 2
    elif partial:
        verdict = '部分完成／待确认：正文、表格、页眉页脚文字与图片均一致，' \
                  '但存在本技能未处理或无法确认的内容，请核对后决定是否交付'
        code = 1
    else:
        verdict = ('完全通过：未发现漏字、增字、漏表、漏图'
                   + ('，也未新增页码' if strict else ''))
        code = 0
    print('\n结论：' + verdict)
    print('注意：本脚本只比对内容，不能证明版式正确，版式仍需逐页查看。')
    return code


if __name__ == '__main__':
    sys.exit(main())
