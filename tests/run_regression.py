#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""official-doc 验收回归测试。

覆盖三类缺陷的修复 + 正常样本，可重复运行：

  缺陷一：段落截断 / 少字 / 多字 / 替换 / 重复段落丢失 / 顺序交换
  缺陷二：页眉页脚文字丢失（含页脚同时有 PAGE 域与文字）
  缺陷三：strict 模式自动编号丢失

用法：
    python3 tests/run_regression.py

退出码：0 全部通过；非 0 有失败（打印失败项）。
"""
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SCRIPTS = os.path.join(ROOT, 'scripts')
VERIFY = os.path.join(SCRIPTS, 'verify_output.py')
INSPECT = os.path.join(SCRIPTS, 'inspect_docx.py')
GENERATE = os.path.join(SCRIPTS, 'generate_gongwen.py')
TESTS = HERE

sys.path.insert(0, SCRIPTS)
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ── 构造样本 ────────────────────────────────────────────

def _add_page_field(para):
    run = para.add_run()
    fb = OxmlElement('w:fldChar'); fb.set(qn('w:fldCharType'), 'begin')
    ins = OxmlElement('w:instrText'); ins.set(qn('xml:space'), 'preserve')
    ins.text = ' PAGE '
    fe = OxmlElement('w:fldChar'); fe.set(qn('w:fldCharType'), 'end')
    r2 = para.add_run(); r2._r.append(fb)
    r3 = para.add_run(); r3._r.append(ins)
    r4 = para.add_run(); r4._r.append(fe)
    para.add_run('8')          # 缓存数字


def make_src_full(path):
    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = '内部资料，请勿外传'
    fp = sec.footer.paragraphs[0]
    fp.text = '机密'
    _add_page_field(fp)
    doc.add_paragraph('关于测试的请示')
    doc.add_paragraph('检测结果不合格，禁止投入使用。')
    doc.add_paragraph('本合同金额为100万元，大写壹佰万元整。')
    doc.save(path)


def make_doc(path, paras, header=None, footer=None, page_field=False):
    doc = Document()
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    if footer:
        fp = doc.sections[0].footer.paragraphs[0]
        fp.text = footer
        if page_field:
            _add_page_field(fp)
    for p in paras:
        doc.add_paragraph(p)
    doc.save(path)


def make_numbered_src(path):
    doc = Document()
    doc.add_paragraph('关于测试的请示')
    p1 = doc.add_paragraph('资产管理工作')
    p2 = doc.add_paragraph('规范账务')
    p3 = doc.add_paragraph('上线公物仓')
    root = doc.part.numbering_part.element
    a = OxmlElement('w:abstractNum'); a.set(qn('w:abstractNumId'), '100')
    lvl0 = OxmlElement('w:lvl'); lvl0.set(qn('w:ilvl'), '0')
    f0 = OxmlElement('w:numFmt'); f0.set(qn('w:val'), 'chineseCounting'); lvl0.append(f0)
    t0 = OxmlElement('w:lvlText'); t0.set(qn('w:val'), '%1、'); lvl0.append(t0)
    s0 = OxmlElement('w:start'); s0.set(qn('w:val'), '1'); lvl0.append(s0)
    a.append(lvl0)
    lvl1 = OxmlElement('w:lvl'); lvl1.set(qn('w:ilvl'), '1')
    f1 = OxmlElement('w:numFmt'); f1.set(qn('w:val'), 'chineseCounting'); lvl1.append(f1)
    t1 = OxmlElement('w:lvlText'); t1.set(qn('w:val'), '（%2）'); lvl1.append(t1)
    s1 = OxmlElement('w:start'); s1.set(qn('w:val'), '1'); lvl1.append(s1)
    a.append(lvl1)
    root.append(a)
    n = OxmlElement('w:num'); n.set(qn('w:numId'), '10')
    aid = OxmlElement('w:abstractNumId'); aid.set(qn('w:val'), '100')
    n.append(aid); root.append(n)

    def add_num(para, ilvl):
        pPr = para._p.get_or_add_pPr()
        numPr = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), str(ilvl))
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), '10')
        numPr.append(il); numPr.append(ni)
        pPr.append(numPr)
    add_num(p1, 0); add_num(p2, 1); add_num(p3, 1)
    doc.save(path)


# ── 运行与断言 ──────────────────────────────────────────

def run_verify(src, dst, meta=None, expect_code=0):
    cmd = [sys.executable, VERIFY, src, dst]
    if meta:
        cmd.append(meta)
    r = subprocess.run(cmd, capture_output=True, text=True)
    return r.returncode, r.stdout


def run_generate(meta, out):
    r = subprocess.run([sys.executable, GENERATE, '--config', meta,
                        '--output', out], capture_output=True, text=True)
    return r.returncode


def check(name, actual, expected):
    ok = actual == expected
    print(('  PASS' if ok else '  FAIL') + f'  {name}: 期望 {expected}，实际 {actual}')
    return ok


def main():
    tmp = tempfile.mkdtemp(prefix='official_doc_reg_')
    results = []

    # ── 正常样本：完整成品（含页眉页脚文字 + 页码）应完全通过 ──
    src_full = os.path.join(tmp, 'src_full.docx')
    dst_full = os.path.join(tmp, 'dst_full.docx')
    make_src_full(src_full)
    make_doc(dst_full,
             ['关于测试的请示', '检测结果不合格，禁止投入使用。',
              '本合同金额为100万元，大写壹佰万元整。'],
             header='内部资料，请勿外传', footer='机密', page_field=True)
    code, _ = run_verify(src_full, dst_full)
    results.append(check('正常样本完全通过', code, 0))

    # ── 缺陷一：截断 ──
    dst_trunc = os.path.join(tmp, 'dst_trunc.docx')
    make_doc(dst_trunc, ['关于测试的请示', '检测结果',
                         '本合同金额为100万元，大写壹佰万元整。'])
    code, out = run_verify(src_full, dst_trunc)
    results.append(check('截断判不通过', code, 2))
    results.append(check('截断报告缺失文字', '不合格，禁止投入使用' in out, True))

    # ── 缺陷一：删除否定词 ──
    dst_neg = os.path.join(tmp, 'dst_neg.docx')
    make_doc(dst_neg, ['关于测试的请示', '检测结果合格，禁止投入使用。',
                       '本合同金额为100万元，大写壹佰万元整。'])
    code, _ = run_verify(src_full, dst_neg)
    results.append(check('删除否定词判不通过', code, 2))

    # ── 缺陷一：修改金额 ──
    dst_amt = os.path.join(tmp, 'dst_amt.docx')
    make_doc(dst_amt, ['关于测试的请示', '检测结果不合格，禁止投入使用。',
                       '本合同金额为1000万元，大写壹佰万元整。'])
    code, _ = run_verify(src_full, dst_amt)
    results.append(check('修改金额判不通过', code, 2))

    # ── 缺陷一：增加前缀 ──
    dst_pre = os.path.join(tmp, 'dst_pre.docx')
    make_doc(dst_pre, ['关于测试的请示', '【新增】检测结果不合格，禁止投入使用。',
                       '本合同金额为100万元，大写壹佰万元整。'])
    code, _ = run_verify(src_full, dst_pre)
    results.append(check('增加前缀判不通过', code, 2))

    # ── 缺陷一：重复段落丢失 ──
    src_dup = os.path.join(tmp, 'src_dup.docx')
    make_doc(src_dup, ['关于测试的请示',
                       '检测结果不合格，禁止投入使用。',
                       '检测结果不合格，禁止投入使用。',
                       '本合同金额为100万元，大写壹佰万元整。'])
    dst_dup = os.path.join(tmp, 'dst_dup.docx')
    make_doc(dst_dup, ['关于测试的请示',
                       '检测结果不合格，禁止投入使用。',
                       '本合同金额为100万元，大写壹佰万元整。'])
    code, _ = run_verify(src_dup, dst_dup)
    results.append(check('重复段落丢失判不通过', code, 2))

    # ── 缺陷一：顺序交换 ──
    dst_swap = os.path.join(tmp, 'dst_swap.docx')
    make_doc(dst_swap, ['关于测试的请示',
                        '本合同金额为100万元，大写壹佰万元整。',
                        '检测结果不合格，禁止投入使用。'],
             header='内部资料，请勿外传', footer='机密', page_field=True)
    code, out = run_verify(src_full, dst_swap)
    results.append(check('顺序交换判不通过', code, 2))
    results.append(check('顺序交换报告', '顺序' in out, True))

    # ── 缺陷二：页眉文字丢失 ──
    dst_nohdr = os.path.join(tmp, 'dst_nohdr.docx')
    make_doc(dst_nohdr, ['关于测试的请示', '检测结果不合格，禁止投入使用。',
                         '本合同金额为100万元，大写壹佰万元整。'])
    code, _ = run_verify(src_full, dst_nohdr)
    results.append(check('页眉文字丢失判不通过', code, 2))

    # ── 缺陷二：页脚文字 + PAGE 域，文字删除 ──
    dst_nofooter = os.path.join(tmp, 'dst_nofooter.docx')
    make_doc(dst_nofooter, ['关于测试的请示', '检测结果不合格，禁止投入使用。',
                            '本合同金额为100万元，大写壹佰万元整。'],
             header='内部资料，请勿外传', footer='', page_field=True)
    code, _ = run_verify(src_full, dst_nofooter)
    results.append(check('页脚文字（含PAGE域）删除判不通过', code, 2))

    # ── 缺陷三：自动编号保留 ──
    src_num = os.path.join(tmp, 'src_num.docx')
    make_numbered_src(src_num)
    meta_num = os.path.join(tmp, 'src_num.json')
    subprocess.run([sys.executable, INSPECT, src_num, '-o', meta_num,
                    '--fidelity', 'strict'], capture_output=True, text=True)
    out_num = os.path.join(tmp, 'out_num.docx')
    run_generate(meta_num, out_num)
    code, _ = run_verify(src_num, out_num, meta_num)
    results.append(check('自动编号保留完全通过', code, 0))

    # ── 缺陷三：自动编号丢失 ──
    dst_nonum = os.path.join(tmp, 'dst_nonum.docx')
    d = Document(out_num)
    for p in d.paragraphs:
        pPr = p._p.pPr
        if pPr is not None:
            numPr = pPr.find(qn('w:numPr'))
            if numPr is not None:
                pPr.remove(numPr)
    d.save(dst_nonum)
    code, _ = run_verify(src_num, dst_nonum, meta_num)
    results.append(check('自动编号丢失判不通过', code, 2))

    passed = sum(results)
    total = len(results)
    print(f'\n===== 回归测试结果：{passed}/{total} 通过 =====')
    return 0 if passed == total else 1


if __name__ == '__main__':
    sys.exit(main())
